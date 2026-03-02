from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from .ollama_client import OllamaTranslator


@dataclass
class TranslateConfig:
    model: str
    output_dir: Path
    base_url: str = "http://localhost:11434"


def _iter_shapes_recursive(shapes) -> Iterable:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes_recursive(shape.shapes)


def _is_probably_japanese(text: str) -> bool:
    if not text.strip():
        return True
    jp = sum(1 for ch in text if ("\u3040" <= ch <= "\u30ff") or ("\u4e00" <= ch <= "\u9fff"))
    return jp / max(1, len(text)) > 0.25


def _set_font_preserve_size(run, fallback_pt: float = 18.0) -> None:
    original_pt = run.font.size.pt if run.font.size else fallback_pt
    run.font.name = "Meiryo"
    run.font.size = Pt(original_pt)


def _adjust_short_bullet_width(shape, paragraph, translated_text: str) -> None:
    if paragraph.level <= 0:
        return
    if len(translated_text.strip()) > 25:
        return
    if not paragraph.runs:
        return
    font_pt = paragraph.runs[0].font.size.pt if paragraph.runs[0].font.size else 18.0
    est_width_pt = max(60.0, len(translated_text.strip()) * font_pt * 1.1)
    if shape.width.pt > est_width_pt:
        shape.width = Pt(est_width_pt)


def _translate_paragraph(paragraph, shape, translator: OllamaTranslator) -> None:
    runs = list(paragraph.runs)
    original = "".join(run.text for run in runs) or paragraph.text
    if not original.strip():
        return

    # Repeat-translation guard
    if _is_probably_japanese(original):
        for run in runs:
            _set_font_preserve_size(run)
        return

    short_bullet = paragraph.level > 0 and len(original.strip()) <= 25
    translated = translator.translate_en_to_ja(original, short_bullet=short_bullet)

    if runs:
        runs[0].text = translated
        for run in runs[1:]:
            run.text = ""
        _set_font_preserve_size(runs[0], runs[0].font.size.pt if runs[0].font.size else 18.0)
    else:
        paragraph.text = translated

    _adjust_short_bullet_width(shape, paragraph, translated)


def _translate_text_frame(shape, translator: OllamaTranslator) -> None:
    tf = shape.text_frame
    for paragraph in tf.paragraphs:
        _translate_paragraph(paragraph, shape, translator)


def _translate_table(shape, translator: OllamaTranslator) -> None:
    table = shape.table
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                original = "".join(run.text for run in paragraph.runs) or paragraph.text
                if not original.strip() or _is_probably_japanese(original):
                    for run in paragraph.runs:
                        _set_font_preserve_size(run)
                    continue
                translated = translator.translate_en_to_ja(original, short_bullet=(paragraph.level > 0 and len(original.strip()) <= 25))
                if paragraph.runs:
                    paragraph.runs[0].text = translated
                    for run in paragraph.runs[1:]:
                        run.text = ""
                    _set_font_preserve_size(paragraph.runs[0])
                else:
                    paragraph.text = translated


def translate_ppt(input_path: Path, config: TranslateConfig) -> Path:
    prs = Presentation(str(input_path))
    translator = OllamaTranslator(model=config.model, base_url=config.base_url)

    for slide in prs.slides:
        for shape in _iter_shapes_recursive(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                _translate_text_frame(shape, translator)
            if getattr(shape, "has_table", False) and shape.has_table:
                _translate_table(shape, translator)

    config.output_dir.mkdir(parents=True, exist_ok=True)
    out_path = config.output_dir / f"{input_path.stem}_ja.pptx"
    prs.save(str(out_path))
    return out_path

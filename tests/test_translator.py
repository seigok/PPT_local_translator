from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from ppt_local_translator import translator as t


class DummyTranslator:
    def __init__(self, *args, **kwargs):
        pass

    def translate_en_to_ja(self, text: str, *, short_bullet: bool = False) -> str:
        return f"訳:{text}"


def test_translate_basic_shape(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(t, "OllamaTranslator", DummyTranslator)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Cloud Native Architecture"
    run.font.bold = True
    run.font.size = Pt(20)

    in_path = tmp_path / "in.pptx"
    prs.save(in_path)

    out = t.translate_ppt(in_path, t.TranslateConfig(model="translategemma:4b", output_dir=tmp_path))
    out_prs = Presentation(out)
    out_shape = out_prs.slides[0].shapes[0]
    run2 = out_shape.text_frame.paragraphs[0].runs[0]

    assert out.exists()
    assert run2.text.startswith("訳:")
    assert run2.font.name == "Meiryo"
    assert round(run2.font.size.pt, 1) == 20.0


def test_skip_retranslation_if_japanese(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(t, "OllamaTranslator", DummyTranslator)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box.text_frame.paragraphs[0].text = "これは日本語の文章である"

    in_path = tmp_path / "jp.pptx"
    prs.save(in_path)

    out = t.translate_ppt(in_path, t.TranslateConfig(model="translategemma:4b", output_dir=tmp_path))
    out_prs = Presentation(out)
    text = out_prs.slides[0].shapes[0].text_frame.paragraphs[0].text
    assert text == "これは日本語の文章である"


def test_adjust_short_bullet_width(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(t, "OllamaTranslator", DummyTranslator)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    p = box.text_frame.paragraphs[0]
    p.level = 1
    p.text = "short bullet"

    in_path = tmp_path / "b.pptx"
    prs.save(in_path)

    out = t.translate_ppt(in_path, t.TranslateConfig(model="translategemma:4b", output_dir=tmp_path))
    out_prs = Presentation(out)
    out_shape = out_prs.slides[0].shapes[0]
    assert out_shape.width.pt < box.width.pt


def test_multi_line_bullets_use_longest_line_for_width(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(t, "OllamaTranslator", DummyTranslator)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(3))
    tf = box.text_frame
    p1 = tf.paragraphs[0]
    p1.level = 1
    p1.text = "short"
    p2 = tf.add_paragraph()
    p2.level = 1
    p2.text = "this is longest bullet line"
    p3 = tf.add_paragraph()
    p3.level = 1
    p3.text = "mid"

    in_path = tmp_path / "ml.pptx"
    prs.save(in_path)

    out = t.translate_ppt(in_path, t.TranslateConfig(model="translategemma:4b", output_dir=tmp_path))
    out_prs = Presentation(out)
    out_shape = out_prs.slides[0].shapes[0]
    assert out_shape.width.pt < box.width.pt


def test_translate_slide_notes(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(t, "OllamaTranslator", DummyTranslator)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    notes = slide.notes_slide.notes_text_frame
    notes.paragraphs[0].text = "Release notes in English"

    in_path = tmp_path / "notes.pptx"
    prs.save(in_path)

    out = t.translate_ppt(in_path, t.TranslateConfig(model="translategemma:4b", output_dir=tmp_path))
    out_prs = Presentation(out)
    text = out_prs.slides[0].notes_slide.notes_text_frame.paragraphs[0].text
    assert text.startswith("訳:")


def test_keep_blank_middle_line(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(t, "OllamaTranslator", DummyTranslator)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
    tf = box.text_frame
    tf.paragraphs[0].text = "Line one"
    tf.add_paragraph().text = ""
    tf.add_paragraph().text = "Line three"

    in_path = tmp_path / "blankline.pptx"
    prs.save(in_path)

    out = t.translate_ppt(in_path, t.TranslateConfig(model="translategemma:4b", output_dir=tmp_path))
    out_prs = Presentation(out)
    ps = out_prs.slides[0].shapes[0].text_frame.paragraphs
    assert len(ps) >= 3
    assert ps[1].text == ""

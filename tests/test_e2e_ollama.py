from pathlib import Path
import shutil

import pytest
import requests
from pptx import Presentation
from pptx.util import Inches

from ppt_local_translator import translator as t


def _has_ollama() -> bool:
    return shutil.which("ollama") is not None


def _model_available(model: str, base_url: str = "http://localhost:11434") -> bool:
    try:
        r = requests.get(f"{base_url}/api/tags", timeout=5)
        r.raise_for_status()
        data = r.json()
        names = {m.get("name", "") for m in data.get("models", [])}
        return model in names
    except Exception:
        return False


@pytest.mark.e2e
def test_e2e_translate_with_ollama(tmp_path: Path):
    model = "translategemma:4b"
    if not _has_ollama():
        pytest.skip("ollama not installed")
    if not _model_available(model):
        pytest.skip(f"model not available: {model}")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.text = "Cloud-native systems improve developer productivity."

    in_path = tmp_path / "e2e_input.pptx"
    prs.save(in_path)

    out_path = t.translate_ppt(
        in_path,
        t.TranslateConfig(model=model, output_dir=tmp_path, base_url="http://localhost:11434"),
    )

    assert out_path.exists()

    out_prs = Presentation(out_path)
    p = out_prs.slides[0].shapes[0].text_frame.paragraphs[0]
    translated_text = p.text

    assert translated_text.strip()
    assert translated_text != "Cloud-native systems improve developer productivity."
    assert p.runs[0].font.name == "Meiryo"
